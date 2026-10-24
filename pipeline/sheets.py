"""Read the source worksheets.

GoogleSheetsSource reads with a service account and the read-only scope.
LocalSource reads exported .xlsx/.csv copies, for development on your own machine only.
Both return, per worksheet, a table of the requested columns in sheet order, plus a
`_row` column holding the sheet row number.
"""
from __future__ import annotations

import os
import re
from pathlib import Path

import pandas as pd

READONLY_SCOPE = "https://www.googleapis.com/auth/spreadsheets.readonly"
API = "https://sheets.googleapis.com/v4/spreadsheets"


class SourceError(RuntimeError):
    """A worksheet could not be read, or lacks a required column."""


def _col_letter(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _q(ws: str) -> str:
    return "'" + ws.replace("'", "''") + "'"


def header_index(header: list, wanted: list[str], where: str) -> dict[str, int]:
    """Map each wanted header to its first column position; fail if absent or repeated."""
    clean = [("" if h is None else str(h)).strip() for h in header]
    out, missing, repeated = {}, [], []
    for w in wanted:
        hits = [i for i, h in enumerate(clean) if h == w]
        if not hits:
            missing.append(w)
        elif len(hits) > 1:
            repeated.append(w)
        else:
            out[w] = hits[0]
    if missing or repeated:
        msg = []
        if missing:
            msg.append(f"missing column(s) {missing}")
        if repeated:
            msg.append(f"column name(s) used more than once {repeated}")
        raise SourceError(f"{where}: " + "; ".join(msg))
    return out


class GoogleSheetsSource:
    """Reads Google Sheets with refreshable, keyless, read-only credentials.

    In GitHub Actions, google-github-actions/auth writes a Workload Identity Federation credentials file
    (no private key) and points GOOGLE_APPLICATION_CREDENTIALS at it. google-auth loads it here and
    refreshes the access token whenever it expires or Google answers 401, so a long run never reuses an
    expired token. Only the spreadsheets.readonly scope is requested.

    Every request is logged with its attempt number, HTTP status and duration. Google's own error status
    and message are logged with long digit runs masked; tokens and cell values are never logged.
    """

    RETRYABLE = frozenset({429, 500, 502, 503, 504})
    QUOTA_REASONS = frozenset({"rateLimitExceeded", "userRateLimitExceeded", "RATE_LIMIT_EXCEEDED"})
    TIMEOUT = (15, 240)            # connect, read (s). Google gives up on its side at about 180 s.
    BACKOFF_BASE = 10.0            # first retry waits about 5-10 s, doubling each time
    BACKOFF_CAP = 120.0
    MAX_ATTEMPTS = 6               # per request
    SOURCE_BUDGET = 20 * 60        # per source (s): stop retrying one sheet after this
    JOB_BUDGET = 40 * 60           # all sources together (s)

    def __init__(self, credentials=None, session=None, clock=None, sleep=None, rand=None):
        import random
        import time

        self.clock = clock or time.monotonic
        self.sleep = sleep or time.sleep
        self.rand = rand or random.random
        self.account = os.environ.get("GOOGLE_SERVICE_ACCOUNT", "the dashboard service account")
        if session is not None:
            self.session = session
        else:
            from google.auth.transport.requests import AuthorizedSession

            if credentials is None:
                import google.auth
                from google.auth.exceptions import DefaultCredentialsError

                try:
                    credentials, _ = google.auth.default(scopes=[READONLY_SCOPE])
                except DefaultCredentialsError as e:
                    raise SourceError("Google credentials not found: the keyless sign-in step did not run or failed") from e
            # AuthorizedSession refreshes before a request when the token has expired, and once more
            # if Google answers 401; it never sends the token anywhere but Google.
            self.session = AuthorizedSession(credentials, max_refresh_attempts=1)
        self.job_deadline = self.clock() + self.JOB_BUDGET
        self.source_deadline = self.job_deadline

    def start_source(self) -> None:
        """Begin the time budget for the next source (bounded by the job budget)."""
        self.source_deadline = min(self.clock() + self.SOURCE_BUDGET, self.job_deadline)

    @staticmethod
    def _google_error(r) -> tuple[str, str]:
        try:
            err = r.json().get("error", {})
        except ValueError:
            return "", ""
        reasons = [d.get("reason", "") for d in err.get("errors", []) if isinstance(d, dict)]
        reasons += [d.get("reason", "") for d in err.get("details", []) if isinstance(d, dict)]
        msg = re.sub(r"\d{5,}", "#", str(err.get("message", "")))[:120]
        return str(err.get("status", "")), (msg + (f" [{', '.join(x for x in reasons if x)}]" if any(reasons) else ""))

    def _wait_before(self, attempt: int, retry_after: str | None) -> float:
        wait = min(self.BACKOFF_CAP, self.BACKOFF_BASE * 2 ** (attempt - 2)) * (0.5 + self.rand() / 2)
        if retry_after:
            try:
                wait = max(wait, min(float(retry_after), self.BACKOFF_CAP))
            except ValueError:
                pass
        return wait

    def _get(self, url: str, params: dict, where: str = "Google Sheets") -> dict:
        import requests
        from google.auth.exceptions import RefreshError, TransportError

        last = "no attempt"
        retry_after = None
        for attempt in range(1, self.MAX_ATTEMPTS + 1):
            if attempt > 1:
                wait = self._wait_before(attempt, retry_after)
                # Do not start an attempt that cannot finish inside the budget.
                if self.clock() + wait + self.TIMEOUT[1] > self.source_deadline:
                    raise SourceError(f"{where}: Google Sheets still unavailable when the time budget ran out "
                                      f"after {attempt - 1} attempt(s); last: {last}")
                print(f"  {where}: retrying in {wait:.0f}s", flush=True)
                self.sleep(wait)
            t0 = self.clock()
            retry_after = None
            try:
                r = self.session.get(url, params=params, timeout=self.TIMEOUT)
            except RefreshError as e:
                raise SourceError(f"{where}: Google sign-in could not be refreshed ({type(e).__name__}); check the keyless sign-in setup") from None
            except (requests.exceptions.Timeout, requests.exceptions.ConnectionError,
                    requests.exceptions.ChunkedEncodingError, TransportError) as e:
                last = f"no answer ({type(e).__name__}) after {self.clock() - t0:.0f}s"
                print(f"  {where}: attempt {attempt}: {last}", flush=True)
                continue
            dur = self.clock() - t0
            status, detail = ("", "") if r.status_code < 400 else self._google_error(r)
            print(f"  {where}: attempt {attempt}: HTTP {r.status_code} in {dur:.1f}s" + (f" ({status}: {detail})" if status or detail else ""), flush=True)
            last = f"HTTP {r.status_code} {status}".strip()
            quota = r.status_code == 403 and any(q in detail for q in self.QUOTA_REASONS)
            if r.status_code in self.RETRYABLE or quota:
                retry_after = r.headers.get("Retry-After")
                continue
            if r.status_code == 401:
                raise SourceError(f"{where}: Google refused the sign-in (401) even after refreshing it; check the keyless sign-in setup")
            if r.status_code == 403:
                raise SourceError(f"{where}: permission denied: share the spreadsheet with {self.account} as Viewer")
            if r.status_code == 404:
                raise SourceError(f"{where}: spreadsheet or worksheet not found")
            if r.status_code >= 400:
                raise SourceError(f"{where}: Google Sheets API error {r.status_code} {status}".strip())
            try:
                return r.json()
            except ValueError:
                last = "HTTP 200 with an unreadable body"
                print(f"  {where}: attempt {attempt}: {last}", flush=True)
                continue
        raise SourceError(f"{where}: Google Sheets still unavailable after {self.MAX_ATTEMPTS} attempts; last: {last}")

    def read(self, spreadsheet_id: str, worksheet: str, columns: list[str], whole_tab: bool = False) -> pd.DataFrame:
        where = f"{worksheet}"
        if whole_tab:
            # One request for header and rows, in sheet order: for small tabs on slow spreadsheets,
            # where each request costs minutes no matter how small its range is.
            body = self._get(f"{API}/{spreadsheet_id}/values/{_q(worksheet)}",
                             {"valueRenderOption": "UNFORMATTED_VALUE", "dateTimeRenderOption": "SERIAL_NUMBER",
                              "majorDimension": "ROWS"}, where)
            if not str(body.get("range", "")).startswith(_q(worksheet) + "!") and not str(body.get("range", "")).startswith(worksheet + "!"):
                raise SourceError(f"{where}: Google returned a different range than requested; not using it")
            rows = body.get("values")
            if not rows:
                raise SourceError(f"{where}: Google returned no header row; the tab is empty or the read was incomplete")
            idx = header_index(rows[0], columns, where)
            data_rows = rows[1:]
            data = {name: [(r[i] if i < len(r) else None) for r in data_rows] for name, i in idx.items()}
            df = pd.DataFrame(data, columns=list(idx.keys()))
            df["_row"] = range(2, len(data_rows) + 2)
            return df
        head = self._get(f"{API}/{spreadsheet_id}/values/{_q(worksheet)}!1:1", {"valueRenderOption": "UNFORMATTED_VALUE"}, where)
        header = (head.get("values") or [[]])[0]
        idx = header_index(header, columns, where)
        ranges = [f"{_q(worksheet)}!{_col_letter(i)}2:{_col_letter(i)}" for i in idx.values()]
        body = self._get(
            f"{API}/{spreadsheet_id}/values:batchGet",
            {"ranges": ranges, "majorDimension": "COLUMNS", "valueRenderOption": "UNFORMATTED_VALUE", "dateTimeRenderOption": "SERIAL_NUMBER"},
            where,
        )
        vrs = body.get("valueRanges")
        if not isinstance(vrs, list) or len(vrs) != len(ranges):
            raise SourceError(f"{where}: Google returned {0 if not isinstance(vrs, list) else len(vrs)} of {len(ranges)} requested columns; not using a partial read")
        cols = [((vr.get("values") or [[]])[0]) for vr in vrs]
        n = max((len(c) for c in cols), default=0)
        data = {name: (c + [None] * (n - len(c))) for name, c in zip(idx.keys(), cols)}
        df = pd.DataFrame(data, columns=list(idx.keys()))
        df["_row"] = range(2, n + 2)
        return df

    def count_rows(self, spreadsheet_id: str, worksheet: str, has_header: bool) -> int:
        body = self._get(f"{API}/{spreadsheet_id}/values/{_q(worksheet)}!A:E", {"valueRenderOption": "UNFORMATTED_VALUE"}, worksheet)
        rows = body.get("values", [])
        if has_header:
            rows = rows[1:]
        return sum(1 for r in rows if any(v not in (None, "") for v in r[:5]))


class LocalSource:
    """Development only: reads exported copies listed in a JSON map {source_key: path}.

    Files hold real client data: keep them outside the repository.
    """

    def __init__(self, files: dict[str, str]):
        self.files = files

    def _frame(self, path: str, worksheet: str) -> list[list]:
        p = Path(path)
        if p.suffix.lower() == ".csv":
            d = pd.read_csv(p, dtype=str, keep_default_na=False, header=None)
            return d.values.tolist()
        import openpyxl

        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        if worksheet not in wb.sheetnames:
            raise SourceError(f"{worksheet}: worksheet not found in {p.name}")
        return [list(r) for r in wb[worksheet].iter_rows(values_only=True)]

    def read_key(self, key: str, worksheet: str, columns: list[str]) -> pd.DataFrame:
        if key not in self.files:
            raise SourceError(f"{key}: no local file configured")
        rows = self._frame(self.files[key], worksheet)
        if not rows:
            raise SourceError(f"{worksheet}: empty")
        idx = header_index(rows[0], columns, worksheet)
        body = rows[1:]
        while body and all(v in (None, "") for v in body[-1]):
            body.pop()
        data = {name: [(r[i] if i < len(r) else None) for r in body] for name, i in idx.items()}
        df = pd.DataFrame(data)
        df["_row"] = range(2, len(body) + 2)
        return df

    def count_rows_key(self, key: str, worksheet: str, has_header: bool) -> int:
        rows = self._frame(self.files[key + ":closed"] if key + ":closed" in self.files else self.files[key], worksheet)
        if has_header:
            rows = rows[1:]
        return sum(1 for r in rows if any(v not in (None, "") for v in list(r)[:5]))


def read_all(cfg: dict, source, pg_source=None, extra_columns: dict | None = None) -> tuple[dict[str, pd.DataFrame], dict[str, int]]:
    """Read every configured worksheet and each lender's Closed tab row count.

    `extra_columns` ({source key: [column, ...]}) adds read-only columns to a source's read, e.g. the lender
    ownership fields for the closing-agent preview; without it every source is read exactly as configured.

    With sources.ringcx.backend: postgres, RingCX comes from `pg_source` (read-only database) instead of
    the sheet; with compare_with_sheets the sheet is also read, as tables["ringcx_sheet"], for comparison.
    """
    tables, closed_counts = {}, {}
    for key, s in cfg["sources"].items():
        if key == "ringcx" and s.get("backend") == "postgres":
            print("reading ringcx / PostgreSQL (read-only)", flush=True)
            if pg_source is None:
                raise SourceError("ringcx: backend is postgres but no database source was set up")
            tables[key] = pg_source.read()
            if not (s.get("postgres") or {}).get("compare_with_sheets"):
                continue
            key = "ringcx_sheet"
        cols = []
        for v in s["columns"].values():
            if v is None:
                continue
            cols.extend(v if isinstance(v, list) else [v])
        cols = list(dict.fromkeys(cols + list((extra_columns or {}).get(key, []))))
        print(f"reading {key} / {s['worksheet']} ({len(cols)} columns)", flush=True)
        if hasattr(source, "start_source"):
            source.start_source()
        if isinstance(source, LocalSource):
            tables[key] = source.read_key("ringcx" if key == "ringcx_sheet" else key, s["worksheet"], cols)
        else:
            tables[key] = source.read(s["spreadsheet_id"], s["worksheet"], cols, whole_tab=bool(s.get("read_whole_tab")))
        print(f"  {key}: {len(tables[key])} rows", flush=True)
        ct = s.get("closed_tab")
        if ct and cfg["validation"].get("closed_tab_cross_check"):
            if isinstance(source, LocalSource):
                closed_counts[key] = source.count_rows_key(key, ct["worksheet"], ct["has_header"])
            else:
                closed_counts[key] = source.count_rows(s["spreadsheet_id"], ct["worksheet"], ct["has_header"])
    return tables, closed_counts


_SECRETISH = re.compile(r"-----BEGIN [A-Z ]*PRIVATE KEY-----|\"private_key\"\s*:")


def assert_no_secret(text: str, where: str) -> None:
    if _SECRETISH.search(text):
        raise SourceError(f"{where}: looks like it contains a private key; refusing to write it")
